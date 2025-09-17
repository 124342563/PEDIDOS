from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import Product, Customer, Order, OrderItem
from .forms import ProductForm, CustomerForm
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
import openpyxl
from openpyxl.utils import get_column_letter
import json
from django.db import transaction
from django.views.decorators.csrf import csrf_exempt
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch

def export_orders_to_pdf(request, orders, title):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{title}.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter

    p.setFont("Helvetica-Bold", 16)
    p.drawString(inch, height - inch, title)

    p.setFont("Helvetica", 12)
    y = height - 1.5 * inch

    headers = ["# Orden", "Cliente", "Fecha", "Estado"]
    col_widths = [1.5*inch, 2.5*inch, 1.5*inch, 1.5*inch]

    for i, header in enumerate(headers):
        p.drawString(inch + sum(col_widths[:i]), y, header)

    y -= 0.25 * inch

    for order in orders:
        if y < inch: # New page if content gets too low
            p.showPage()
            p.setFont("Helvetica-Bold", 16)
            p.drawString(inch, height - inch, title)
            p.setFont("Helvetica", 12)
            y = height - 1.5 * inch

        data = [
            order.order_number,
            order.customer.name,
            order.created_at.strftime("%Y-%m-%d"),
            order.get_status_display()
        ]
        for i, item in enumerate(data):
            p.drawString(inch + sum(col_widths[:i]), y, item)
        y -= 0.25 * inch


    p.showPage()
    p.save()
    return response

@login_required
def home(request):
    return render(request, 'dashboard.html')

@login_required
def product_list(request):
    products = Product.objects.all().order_by('name')

    # Search logic
    search_query = request.GET.get('search', '')
    if search_query:
        products = products.filter(
            Q(code__icontains=search_query) |
            Q(name__icontains=search_query) |
            Q(composition__icontains=search_query)
        )

    # Filter logic
    line_filter = request.GET.get('line', '')
    if line_filter:
        products = products.filter(line=line_filter)

    status_filter = request.GET.get('status', '')
    if status_filter:
        products = products.filter(status=status_filter)

    context = {
        'products': products,
        'lines': Product.Line.choices,
        'statuses': Product.Status.choices,
        'selected_line': line_filter,
        'selected_status': status_filter,
        'search_query': search_query
    }
    return render(request, 'products.html', context)

@login_required
def customer_list(request):
    customers = Customer.objects.all().order_by('name')

    # Search logic
    search_query = request.GET.get('search', '')
    if search_query:
        customers = customers.filter(
            Q(name__icontains=search_query) |
            Q(nit__icontains=search_query) |
            Q(city__icontains=search_query)
        )

    # Filter logic
    type_filter = request.GET.get('type', '')
    if type_filter:
        customers = customers.filter(customer_type=type_filter)

    context = {
        'customers': customers,
        'customer_types': Customer.CustomerType.choices,
        'selected_type': type_filter,
        'search_query': search_query,
    }
    return render(request, 'customers.html', context)

@csrf_exempt
@login_required
def create_order(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            customer_id = data.get('customer_id')
            items = data.get('items')

            if not customer_id or not items:
                return JsonResponse({'error': 'Faltan datos del cliente o de los productos.'}, status=400)

            with transaction.atomic():
                customer = Customer.objects.get(id=customer_id)

                order = Order.objects.create(
                    customer=customer,
                    created_by=request.user,
                    comments=data.get('comments', '')
                )

                for item_data in items:
                    product = Product.objects.get(id=item_data['product_id'])
                    OrderItem.objects.create(
                        order=order,
                        product=product,
                        quantity=item_data['quantity'],
                        price=product.price
                    )

            return JsonResponse({'success': True, 'redirect_url': '/logistics/'})

        except Customer.DoesNotExist:
            return JsonResponse({'error': 'El cliente no existe.'}, status=400)
        except Product.DoesNotExist:
            return JsonResponse({'error': 'Uno de los productos no existe.'}, status=400)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)

    customers = Customer.objects.all().order_by('name')
    products = Product.objects.filter(status='DISPONIBLE').order_by('name')
    lines = Product.Line.choices

    context = {
        'customers': customers,
        'products': products,
        'lines': lines,
    }
    return render(request, 'create_order.html', context)


@login_required
def order_list_terminado(request):
    orders = Order.objects.filter(items__product__line='TERMINADO').distinct().order_by('-created_at')

    search_query = request.GET.get('search', '')
    if search_query:
        orders = orders.filter(Q(order_number__icontains=search_query) | Q(customer__name__icontains=search_query))

    status_filter = request.GET.get('status', '')
    if status_filter:
        orders = orders.filter(status=status_filter)

    if request.GET.get('format') == 'pdf':
        return export_orders_to_pdf(request, orders, 'Ordenes de Producto Terminado')

    context = {
        'orders': orders,
        'title': 'Órdenes de Producto Terminado',
        'statuses': Order.OrderStatus.choices,
        'selected_status': status_filter,
        'search_query': search_query,
    }
    return render(request, 'order_list.html', context)


# Customer CRUD
@login_required
def customer_create(request):
    if request.method == 'POST':
        form = CustomerForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:customer_list')
    else:
        form = CustomerForm()
    context = {
        'form': form,
        'title': 'Nuevo Cliente',
        'subtitle': 'Añadir un nuevo cliente a la base de datos.'
    }
    return render(request, 'customer_form.html', context)

@login_required
def customer_edit(request, pk):
    customer = get_object_or_404(Customer, pk=pk)
    if request.method == 'POST':
        form = CustomerForm(request.POST, instance=customer)
        if form.is_valid():
            form.save()
            return redirect('core:customer_list')
    else:
        form = CustomerForm(instance=customer)
    context = {
        'form': form,
        'title': 'Editar Cliente',
        'subtitle': f'Editando {customer.name}'
    }
    return render(request, 'customer_form.html', context)

@login_required
def customer_delete(request, pk):
    customer = get_object_or_404(Customer, pk=pk)
    if request.method == 'POST':
        customer.delete()
        return redirect('core:customer_list')
    context = {
        'object': customer,
        'title': 'Confirmar Eliminación',
        'confirm_message': f'¿Estás seguro de que quieres eliminar el cliente "{customer.name}"?'
    }
    return render(request, 'confirm_delete.html', context)


# Product CRUD
@login_required
def product_create(request):
    if request.method == 'POST':
        form = ProductForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:product_list')
    else:
        form = ProductForm()
    context = {
        'form': form,
        'title': 'Nuevo Producto',
        'subtitle': 'Añadir un nuevo producto al catálogo.'
    }
    return render(request, 'product_form.html', context)

@login_required
def product_edit(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            form.save()
            return redirect('core:product_list')
    else:
        form = ProductForm(instance=product)
    context = {
        'form': form,
        'title': 'Editar Producto',
        'subtitle': f'Editando {product.name}'
    }
    return render(request, 'product_form.html', context)

@login_required
def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        product.delete()
        return redirect('core:product_list')
    context = {
        'object': product,
        'title': 'Confirmar Eliminación',
        'confirm_message': f'¿Estás seguro de que quieres eliminar el producto "{product.name}"?'
    }
    return render(request, 'confirm_delete.html', context)

@login_required
def export_logistics_to_excel(request):
    orders = Order.objects.prefetch_related('items', 'items__product', 'customer', 'created_by').all().order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logística de Pedidos"

    headers = [
        "Número de Pedido", "Fecha", "Cliente", "NIT Cliente", "Vendedor",
        "Código Producto", "Producto", "Cantidad", "Precio Unitario", "Precio Total Item",
        "Estado Pedido"
    ]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header_title
        ws.column_dimensions[get_column_letter(col_num)].width = 20

    row_num = 2
    for order in orders:
        for item in order.items.all():
            ws.cell(row=row_num, column=1).value = order.order_number
            ws.cell(row=row_num, column=2).value = order.created_at.strftime("%Y-%m-%d %H:%M")
            ws.cell(row=row_num, column=3).value = order.customer.name
            ws.cell(row=row_num, column=4).value = order.customer.nit
            ws.cell(row=row_num, column=5).value = order.created_by.username
            ws.cell(row=row_num, column=6).value = item.product.code
            ws.cell(row=row_num, column=7).value = item.product.name
            ws.cell(row=row_num, column=8).value = item.quantity
            ws.cell(row=row_num, column=9).value = item.price
            ws.cell(row=row_num, column=10).value = item.quantity * item.price
            ws.cell(row=row_num, column=11).value = order.get_status_display()
            row_num += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="logistica_pedidos.xlsx"'
    wb.save(response)

    return response


@login_required
def logistics_list(request):
    orders = Order.objects.all().order_by('-created_at')

    # Add filtering logic here later if needed

    context = {
        'orders': orders,
        'statuses': Order.OrderStatus.choices,
    }
    return render(request, 'logistics.html', context)


@login_required
def order_list_magistrales(request):
    orders = Order.objects.filter(items__product__line='MAGISTRAL').distinct().order_by('-created_at')

    search_query = request.GET.get('search', '')
    if search_query:
        orders = orders.filter(Q(order_number__icontains=search_query) | Q(customer__name__icontains=search_query))

    status_filter = request.GET.get('status', '')
    if status_filter:
        orders = orders.filter(status=status_filter)

    if request.GET.get('format') == 'pdf':
        return export_orders_to_pdf(request, orders, 'Ordenes Magistrales')

    context = {
        'orders': orders,
        'title': 'Órdenes Magistrales',
        'statuses': Order.OrderStatus.choices,
        'selected_status': status_filter,
        'search_query': search_query,
    }
    return render(request, 'order_list.html', context)
