import os
import uuid
from datetime import datetime, date
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import current_app
from app.extensions import db
from app.models import Colaborador, Entrega, TrackingEvento


def importar_colaboradores_excel(stream, id_cliente):
    df = pd.read_excel(stream, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    creados = 0
    errores = 0
    for _, row in df.iterrows():
        try:
            nombre = row.get("NombreDestinatarioCompleto") or row.get("Nombre") or row.get("nombre")
            if not nombre or str(nombre).strip() in ("nan", ""):
                errores += 1
                continue
            fecha_nac = None
            fn_raw = row.get("FechaNacimiento") or row.get("fecha_nacimiento")
            if fn_raw and str(fn_raw) != "nan":
                try:
                    fecha_nac = pd.to_datetime(fn_raw).date()
                except Exception:
                    pass
            colaborador = Colaborador(
                id_cliente=id_cliente,
                nombre=str(nombre).strip(),
                documento=_safe(row.get("Documento") or row.get("documento")),
                telefono=_safe(row.get("Telefono") or row.get("telefono")),
                correo=_safe(row.get("Correo") or row.get("correo")),
                direccion=_safe(row.get("DireccionEntrega") or row.get("Direccion") or row.get("direccion")),
                ciudad=_safe(row.get("Ciudad") or row.get("ciudad")),
                municipio=_safe(row.get("Municipio") or row.get("municipio")),
                fecha_nacimiento=fecha_nac,
                activo=True,
            )
            db.session.add(colaborador)
            creados += 1
        except Exception:
            errores += 1
    db.session.commit()
    return {"creados": creados, "errores": errores}


def cargar_entregas_desde_base(stream, pedido):
    df = pd.read_excel(stream, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    creadas = 0
    errores = 0

    for _, row in df.iterrows():
        try:
            nombre = _safe(row.get("NombreDestinatarioCompleto") or row.get("Nombre"))
            if not nombre:
                errores += 1
                continue
            fecha_entrega = None
            fe_raw = row.get("FechaEntrega") or row.get("fecha_entrega")
            if pedido.tipo_operacion == "CONVENIO_CUMPLEAÑOS":
                fn_raw = row.get("FechaNacimiento") or row.get("fecha_nacimiento")
                if fn_raw and str(fn_raw) != "nan":
                    try:
                        fn = pd.to_datetime(fn_raw).date()
                        year_actual = date.today().year
                        fecha_entrega = date(year_actual, fn.month, fn.day)
                    except Exception:
                        pass
            elif fe_raw and str(fe_raw) != "nan":
                try:
                    fecha_entrega = pd.to_datetime(fe_raw).date()
                except Exception:
                    fecha_entrega = pedido.fecha_entrega_general

            entrega = Entrega(
                id_pedido=pedido.id_pedido,
                id_cliente=pedido.id_cliente,
                destinatario=nombre,
                empresa=pedido.cliente.nombre if pedido.cliente else None,
                direccion=_safe(row.get("DireccionEntrega") or row.get("Direccion")),
                ciudad=_safe(row.get("Ciudad")),
                municipio=_safe(row.get("Municipio")),
                telefono=_safe(row.get("Telefono")),
                producto=pedido.producto_principal,
                cantidad=int(_safe(row.get("Cantidad")) or 1),
                fecha_entrega=fecha_entrega or pedido.fecha_entrega_general,
                estado_entrega="PENDIENTE",
                observacion=_safe(row.get("ObservacionAdicional") or row.get("Observacion")),
                token_tracking=str(uuid.uuid4()),
            )
            db.session.add(entrega)
            evento = TrackingEvento(
                id_pedido=pedido.id_pedido,
                estado_anterior=None,
                estado_nuevo="PENDIENTE",
                descripcion=f"Entrega creada desde base para {nombre}",
                usuario="sistema",
                origen="SISTEMA",
            )
            db.session.add(evento)
            creadas += 1
        except Exception:
            errores += 1
    db.session.commit()
    return {"creadas": creadas, "errores": errores}


def exportar_plano_simpliroute(pedido):
    export_dir = os.path.join(current_app.config["EXPORT_FOLDER"], "simpliroute")
    os.makedirs(export_dir, exist_ok=True)
    filename = f"simpliroute_{pedido.numero_pedido}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(export_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SimpliRoute"
    headers = ["title", "address", "city", "contact_name", "contact_phone", "reference", "notes", "planned_date"]
    ws.append(headers)
    for entrega in pedido.entregas:
        ws.append([
            entrega.destinatario or "",
            entrega.direccion or "",
            entrega.ciudad or "",
            entrega.destinatario or "",
            entrega.telefono or "",
            pedido.numero_pedido,
            entrega.observacion or "",
            entrega.fecha_entrega.strftime("%Y-%m-%d") if entrega.fecha_entrega else "",
        ])
    wb.save(filepath)
    return filepath


def exportar_plano_alas(pedido):
    export_dir = os.path.join(current_app.config["EXPORT_FOLDER"], "alas")
    os.makedirs(export_dir, exist_ok=True)
    filename = f"alas_{pedido.numero_pedido}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(export_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alas_SISCORE"
    headers = [
        "Factura/Pedido", "Destinatario", "Dirección", "Ciudad", "Municipio",
        "Departamento", "Teléfono", "Contenido", "Cantidad", "Peso",
        "Valor declarado", "Observación"
    ]
    ws.append(headers)
    for entrega in pedido.entregas:
        ws.append([
            pedido.numero_pedido,
            entrega.destinatario or "",
            entrega.direccion or "",
            entrega.ciudad or "",
            entrega.municipio or "",
            "",
            entrega.telefono or "",
            entrega.producto or "",
            entrega.cantidad or 1,
            "",
            "",
            entrega.observacion or "",
        ])
    wb.save(filepath)
    return filepath


def exportar_plano_pibox(pedido):
    export_dir = os.path.join(current_app.config["EXPORT_FOLDER"], "pibox")
    os.makedirs(export_dir, exist_ok=True)
    filename = f"pibox_{pedido.numero_pedido}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(export_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pibox"
    headers = [
        "NumeroPedido", "Destinatario", "Direccion", "Ciudad", "Municipio",
        "Telefono", "Producto", "Cantidad", "FechaEntrega", "Observacion"
    ]
    ws.append(headers)
    for entrega in pedido.entregas:
        ws.append([
            pedido.numero_pedido,
            entrega.destinatario or "",
            entrega.direccion or "",
            entrega.ciudad or "",
            entrega.municipio or "",
            entrega.telefono or "",
            entrega.producto or "",
            entrega.cantidad or 1,
            entrega.fecha_entrega.strftime("%Y-%m-%d") if entrega.fecha_entrega else "",
            entrega.observacion or "",
        ])
    wb.save(filepath)
    return filepath


def exportar_reporte_entregas(pedido):
    export_dir = os.path.join(current_app.config["EXPORT_FOLDER"], "reportes")
    os.makedirs(export_dir, exist_ok=True)
    filename = f"reporte_entregas_{pedido.numero_pedido}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(export_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entregas"
    headers = [
        "ID Entrega", "Destinatario", "Empresa", "Dirección", "Ciudad", "Municipio",
        "Teléfono", "Producto", "Cantidad", "Fecha Entrega", "Estado", "Transportadora",
        "Guía Alas", "Observación"
    ]
    ws.append(headers)
    for entrega in pedido.entregas:
        ws.append([
            entrega.id_entrega,
            entrega.destinatario or "",
            entrega.empresa or "",
            entrega.direccion or "",
            entrega.ciudad or "",
            entrega.municipio or "",
            entrega.telefono or "",
            entrega.producto or "",
            entrega.cantidad or 1,
            entrega.fecha_entrega.strftime("%Y-%m-%d") if entrega.fecha_entrega else "",
            entrega.estado_entrega or "",
            entrega.transportadora or "",
            entrega.guia_alas or "",
            entrega.observacion or "",
        ])
    wb.save(filepath)
    return filepath


def generar_plantilla_cargue():
    export_dir = os.path.join(current_app.config["EXPORT_FOLDER"], "generales")
    os.makedirs(export_dir, exist_ok=True)
    filepath = os.path.join(export_dir, "plantilla_cargue_despachos.xlsx")

    wb = openpyxl.Workbook()

    # Sheet 1: INSTRUCCIONES
    ws_inst = wb.active
    ws_inst.title = "INSTRUCCIONES"
    ws_inst["A1"] = "PLANTILLA DE CARGUE DE DESPACHOS - TGC LOGÍSTICA"
    ws_inst["A1"].font = Font(bold=True, size=14)
    instrucciones = [
        ("A3", "INSTRUCCIONES GENERALES"),
        ("A4", "1. Complete la hoja CARGUE_DESPACHOS con los datos de cada entrega."),
        ("A5", "2. Los campos marcados con (*) son obligatorios."),
        ("A6", "3. Use la hoja LISTAS para ver los valores permitidos en cada campo."),
        ("A7", "4. No modifique los encabezados de las columnas."),
        ("A8", "5. Para pedidos con varios productos, use la hoja DETALLE_ITEMS."),
        ("A10", "TIPOS DE OPERACIÓN VÁLIDOS:"),
        ("A11", "NORMAL | CONVENIO_CUMPLEAÑOS | MASIVO_1_A_1 | UN_SOLO_PUNTO | RECOGE_EN_BODEGA"),
        ("A13", "TIPOS DE DESPACHO VÁLIDOS:"),
        ("A14", "TOTAL | PARCIAL | FINAL"),
    ]
    for cell, val in instrucciones:
        ws_inst[cell] = val

    # Sheet 2: CARGUE_DESPACHOS
    ws_cargue = wb.create_sheet("CARGUE_DESPACHOS")
    headers_obligatorios = [
        "NumeroPedido*", "TipoOperacion*", "TipoDespacho*", "ClienteEmpresa*",
        "NombreDestinatarioCompleto*", "DireccionEntrega*", "Telefono*",
        "Ciudad*", "Municipio*", "FechaEntrega*", "Producto*", "Cantidad*",
        "Transportadora*", "EstadoInicial*", "ObservacionAdicional"
    ]
    headers_opcionales = [
        "Documento", "Correo", "ContactoResponsable", "AreaDependencia",
        "ValorDeclarado", "Peso", "GuiaAlas", "LinkSimpliRoute",
        "TrackingPibox", "ReferenciaInterna"
    ]
    all_headers = headers_obligatorios + headers_opcionales
    ws_cargue.append(all_headers)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    opt_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    for i, cell in enumerate(ws_cargue[1], 1):
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill if i <= len(headers_obligatorios) else opt_fill
        cell.alignment = Alignment(horizontal="center")

    # Sheet 3: DETALLE_ITEMS
    ws_items = wb.create_sheet("DETALLE_ITEMS")
    ws_items.append(["NumeroPedido*", "Producto*", "Cantidad*", "Observacion"])

    # Sheet 4: LISTAS
    ws_listas = wb.create_sheet("LISTAS")
    ws_listas["A1"] = "TIPOS DE OPERACIÓN"
    for i, t in enumerate(["NORMAL", "CONVENIO_CUMPLEAÑOS", "MASIVO_1_A_1", "UN_SOLO_PUNTO", "RECOGE_EN_BODEGA"], 2):
        ws_listas.cell(row=i, column=1, value=t)
    ws_listas["C1"] = "TRANSPORTADORAS"
    for i, t in enumerate(["SIMPLIROUTE", "ALAS", "PIBOX", "UBER", "MENSAJEROS URBANOS", "INTERNO", "RECOGE EN BODEGA"], 2):
        ws_listas.cell(row=i, column=3, value=t)
    ws_listas["E1"] = "TIPOS DE DESPACHO"
    for i, t in enumerate(["TOTAL", "PARCIAL", "FINAL"], 2):
        ws_listas.cell(row=i, column=5, value=t)

    # Sheet 5: EJEMPLOS
    ws_ej = wb.create_sheet("EJEMPLOS")
    ws_ej.append(all_headers)
    ws_ej.append([
        "67336", "CONVENIO_CUMPLEAÑOS", "TOTAL", "ZIMMER BIOMET",
        "Andrés Castillo", "Calle 163 #54-20", "3001234567",
        "Bogotá", "Bogotá D.C.", "2026-12-31", "Torta de cumpleaños", "1",
        "ALAS", "PENDIENTE", "Sin novedad",
        "12345678", "andres@email.com", "Recursos Humanos", "Bienestar",
        "50000", "0.5", "", "", "", "INT-001"
    ])

    wb.save(filepath)
    return filepath


def _safe(value):
    if value is None:
        return None
    s = str(value).strip()
    return None if s in ("nan", "", "None") else s
